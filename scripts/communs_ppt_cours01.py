#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Génère la présentation du Cours 01 « Introduction à l'IA » (30 diapositives).

Pourquoi une reconstruction plutôt qu'une retouche : la version en ligne
comporte 11 diapositives et ne parle de l'IA que comme d'un assistant de
conversation. Les agents (Cursor, OpenCode), la réparation et le flashage
d'appareils, la création d'un site web, les astuces de terrain et l'usage
pédagogique FLE en sont absents — alors que ce sont précisément lesquestions
que pose un étudiant de 2ᵉ année PEP qui vient de FLE et qui veut un résultat
immédiat.

La charte, les logos, la calibration des tailles et les schémas viennent de
scripts/pptx_commun.py : la présentation reste visuellement identique à
celle du site, et aux quatre autres cours.

Usage :
    python3 scripts/generer_ppt_cours01.py [chemin_de_sortie.pptx]
"""
import os
import sys

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

from pptx_commun import (AMBER, AMBER_L, BLUE, BLUE_L, DARK, FILET_MARIN,
                         GREEN, GREEN_L, GREY, GREY_L, LARGEUR, LINE_C,
                         NAVY, NAVY_2, STYLES, TAILLE_ARROW, TAILLE_BOITE,
                         TAILLE_PASTILLE, TAILLE_TEXTE_MAX, TAILLE_TITRE,
                         TITRE_CLAIR, WHITE,
                         add_footer, add_paragraph, add_shape, add_textbox,
                         couverture, dbox, draw_flow, ecrire_notes,
                         nouveau_conteneur, taille_qui_tient, tailles_boite)

RACINE = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
IMG_UNIV = os.path.join(RACINE, "images", "Université_de_Mostaganem.png")
IMG_FLE = os.path.join(RACINE, "images", "LOGO-FLE-UNIV-Mosta.jpeg")
DEFAUT = os.path.join(RACINE, "cours", "Module Intelligence Artificielle",
                      "Presentations_PPTX", "QR_IA_01_Introduction_IA.pptx")

PIED = "Q/R — Cours 01 · Introduction à l'IA   |   Dr. Madani BELACEL — Université de Mostaganem"

# Géométrie relevée sur la présentation existante, pour un rendu identique.
Y_EYEBROW, H_EYEBROW = 0.58, 0.34
Y_TITRE = 1.06
Y_FILET = 2.14
Y_CHAPO = 2.24
Y_CHAPO_TXT = 2.38
L_G, L_D = 0.75, 11.83
L_TXT = 0.85
W_TXT = 11.63


# ---------------------------------------------------------------------------
# Charpente d'une diapositive
# ---------------------------------------------------------------------------
def bandeau(slide):
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, LARGEUR, 0.26, fill=NAVY)


def entete(slide, idee, titre, sous_partie=None):
    """Bandeau du haut, pastille « IDÉE n », titre et filet."""
    bandeau(slide)
    if idee:
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, L_G, Y_EYEBROW, 1.35,
                  H_EYEBROW, fill=BLUE_L, line=BLUE, line_w=1.0, radius=0.25)
        tf = add_textbox(slide, L_G, Y_EYEBROW, 1.35, H_EYEBROW,
                         anchor=MSO_ANCHOR.MIDDLE)
        add_paragraph(tf, idee, size=TAILLE_PASTILLE, bold=True, color=NAVY,
                      align=PP_ALIGN.CENTER, space_after=0)
    tf = add_textbox(slide, L_G, Y_TITRE, L_D, 0.70)
    taille = taille_qui_tient(titre, L_D, 0.70, maxi=TAILLE_TITRE, mini=19)
    add_paragraph(tf, titre, size=taille, bold=True, color=NAVY, space_after=0)
    if sous_partie:
        add_shape(slide, MSO_SHAPE.RECTANGLE, L_G, Y_FILET, L_D, 0.01,
                  fill=LINE_C)
        tf = add_textbox(slide, L_G, Y_FILET + 0.04, L_D, 0.26)
        add_paragraph(tf, sous_partie, size=11.5, italic=True, color=GREY,
                      space_after=0)
    else:
        add_shape(slide, MSO_SHAPE.RECTANGLE, L_G, Y_FILET, L_D, 0.01,
                  fill=LINE_C)


def chapo(slide, texte, hauteur=0.95):
    """Le paragraphe d'accroche, sur fond clair."""
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 0.55, Y_CHAPO, 12.23,
              hauteur, fill=GREY_L, line=None, radius=0.1)
    tf = add_textbox(slide, L_TXT, Y_CHAPO + 0.14, W_TXT, hauteur - 0.28)
    taille = taille_qui_tient(texte, W_TXT, hauteur - 0.28, maxi=TAILLE_TEXTE_MAX,
                              mini=12)
    add_paragraph(tf, texte, size=taille, color=DARK, space_after=0,
                  line_spacing=1.16)


def section(slide, numero, titre, chapeau, items, notes=None):
    """Diapositive de séparation de partie."""
    bandeau(slide)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0.26, LARGEUR, 6.36, fill=NAVY)
    add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.1, 2.05, 1.05, 1.05,
              fill=NAVY_2, line=BLUE, line_w=2.0, radius=0.5)
    tf = add_textbox(slide, 1.1, 2.05, 1.05, 1.05, anchor=MSO_ANCHOR.MIDDLE)
    add_paragraph(tf, str(numero), size=30, bold=True, color=WHITE,
                  align=PP_ALIGN.CENTER, space_after=0)
    tf = add_textbox(slide, 2.45, 2.05, 9.4, 1.05, anchor=MSO_ANCHOR.MIDDLE)
    add_paragraph(tf, titre, size=34, bold=True, color=WHITE, space_after=0)
    tf = add_textbox(slide, 1.1, 3.45, 11.1, 1.4)
    add_paragraph(tf, chapeau, size=17, color=WHITE, space_after=10,
                  line_spacing=1.25)
    for it in items:
        add_paragraph(tf, [("→  ", {"color": TITRE_CLAIR}), (it, {})],
                      size=14.5, color=TITRE_CLAIR, space_after=6)
    if notes:
        ecrire_notes(slide, notes)


# ---------------------------------------------------------------------------
# Blocs de contenu
# ---------------------------------------------------------------------------
def ecrire_boite(slide, x, y, w, h, titre, legende, style):
    """Boîte àcoins arrondis, avec son titre et sa légende.

    dbox() se contente de créer la forme et de rendre son cadre de texte :
    c'est à l'appelant de le remplir. Sans ce remplissage, les boîtes
    s'affichent vides au rendu — ce qui est exactement ce qui s'est produit
    sur une première version de cette présentation.
    """
    fill, line_c, txt_c, sub_c = STYLES.get(style, STYLES["primary"])
    sh, tf = dbox(slide, x, y, w, h, titre, legende, style)
    t_grand, t_petit = tailles_boite(w)
    if titre:
        add_paragraph(tf, titre, size=t_grand, bold=True, color=txt_c,
                      align=PP_ALIGN.CENTER, space_after=3, line_spacing=1.08)
    for ligne in (legende or "").split("\n"):
        if ligne.strip():
            add_paragraph(tf, ligne.strip(), size=t_petit, color=sub_c,
                          align=PP_ALIGN.CENTER, space_after=1,
                          line_spacing=1.12)
    return sh


BAS_CONTENU = 6.78        # dernier pixel utilisable, au-dessus du filet de pied


def grille(slide, cases, y, h=None, colonnes=None, x=L_G, w=L_D, bas=BAS_CONTENU):
    """Grille de boites, hauteur calculée pour tenir dans l'espace restant.

    cases : liste de (titre, legende, style) ou (titre, legende, mention,
    style) quand la légende doit être la ligne principale.

    La hauteur se déduit du nombre de rangées et de la limite basse. Sans
    cela, une grille de six cases sur deux rangées débordait sous le pied de
    page — exactement ce que scripts/verifier_pptx.py signale.
    """
    n = colonnes or len(cases)
    gap = 0.18
    bw = (w - gap * (n - 1)) / n
    rangees = (len(cases) + n - 1) // n
    if h is None:
        h = (bas - y - gap * (rangees - 1)) / rangees
    for i, case in enumerate(cases):
        if len(case) == 4:
            titre, legende, mention, style = case
        else:
            titre, legende, style = case
            mention = ""
        col, rang = i % n, i // n
        bx = x + col * (bw + gap)
        by = y + rang * (h + gap)
        ecrire_boite(slide, bx, by, bw, h, titre,
                     "%s\n%s" % (legende, mention) if mention else legende,
                     style)


def deux_colonnes(slide, gauche, droite, y, h, x=L_G, w=L_D, gap=0.3):
    bw = (w - gap) / 2
    ecrire_boite(slide, x, y, bw, h, gauche[0], gauche[1], gauche[2])
    ecrire_boite(slide, x + bw + gap, y, bw, h, droite[0], droite[1],
                 droite[2])


def tableau(slide, entetes, lignes, y, h, largeurs=None, x=L_G, w=L_D):
    """Tableau simple : bandeau d'entete + lignes alternees."""
    n = len(entetes)
    lw = largeurs or [w / n] * n
    row_h = h / (len(lignes) + 1)
    cx = x
    for i, e in enumerate(entetes):
        add_shape(slide, MSO_SHAPE.RECTANGLE, cx, y, lw[i] - 0.03, row_h,
                  fill=NAVY)
        tf = add_textbox(slide, cx + 0.12, y, lw[i] - 0.27, row_h,
                         anchor=MSO_ANCHOR.MIDDLE)
        add_paragraph(tf, e, size=12.5, bold=True, color=WHITE, space_after=0)
        cx += lw[i]
    for r, ligne in enumerate(lignes):
        cx = x
        fond = WHITE if r % 2 else GREY_L
        add_shape(slide, MSO_SHAPE.RECTANGLE, x, y + (r + 1) * row_h, w,
                  row_h - 0.03, fill=fond)
        for i, cell in enumerate(ligne):
            tf = add_textbox(slide, cx + 0.12, y + (r + 1) * row_h,
                             lw[i] - 0.27, row_h - 0.03,
                             anchor=MSO_ANCHOR.MIDDLE)
            taille = taille_qui_tient(cell, lw[i] - 0.27, row_h - 0.06,
                                      maxi=12.5, mini=8.5)
            add_paragraph(tf, cell, size=taille,
                          bold=(i == 0), color=DARK if i else NAVY,
                          space_after=0, line_spacing=1.05)
            cx += lw[i]


def notes(idee, titre, objectif, discours, qr, piege):
    """Notes du présentateur : les quatre blocs demandes."""
    q = "\n\n".join("Q%d. %s\nR%d. %s" % (i + 1, question, i + 1, reponse)
                    for i, (question, reponse) in enumerate(qr))
    return """%s — %s

🎯 OBJECTIF
%s

🗣️ CE QUE JE DIS
%s

── QUESTIONS & RÉPONSES (FR) ──
%s

⚠️ PIÈGE À ÉVITER
%s""" % (idee, titre, objectif, discours, q, piege)