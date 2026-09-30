#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Controle automatise du rendu d'un PPTX.

Repete ce que fait l'oeil sur un LibreOffice converti en PDF :
  - aucune forme ne depasse de la diapositive ;
  - aucune forme ne chevauche le pied de page (sauf la couverture, qui
    n'a pas de pied) ;
  - le texte de chaque boite tient dans sa boite. La largeur moyenne d'un
    caractere est celle calibree dans scripts/pptx_commun.py : 0.65 a 11 pt,
    0.42 a 38 pt, interpolate en logarithme.

Usage : python3 verifier_pptx.py fichier.pptx
Sortie : 0 si tout va bien, 1 sinon, la liste des anomalies sur stdout.
"""
import math
import os
import sys

from pptx import Presentation
from pptx.util import Emu

PIED = 6.92          # filet au-dessus du pied de page
INTER = 1.15        # interligne utilise par les generateurs


def ratio_texte(taille):
    return max(0.40, min(0.70, 1.094 - 0.427 * math.log10(max(1.0, taille))))


def pouces(v):
    return Emu(v).inches if v is not None else None


def main():
    if len(sys.argv) < 2:
        print(__doc__)
        return 2
    chemin = sys.argv[1]
    prs = Presentation(chemin)
    L, H = pouces(prs.slide_width), pouces(prs.slide_height)
    problemes = []

    for n, s in enumerate(prs.slides, 1):
        for sh in s.shapes:
            try:
                x, y = pouces(sh.left), pouces(sh.top)
                w, h = pouces(sh.width), pouces(sh.height)
            except (TypeError, AttributeError):
                continue
            if None in (x, y, w, h):
                continue

            # un tableau PowerPoint n'a pas toujours left/top exploitable
            if x < -0.01 or y < -0.01 or x + w > L + 0.02 or y + h > H + 0.02:
                problemes.append("diapo %d : forme hors cadre (%.2f,%.2f %.2fx%.2f)"
                                % (n, x, y, w, h))
            if n > 1 and y < PIED and y + h > PIED + 0.03 and h > 0.20:
                txt = (sh.text_frame.text.strip()[:28] if sh.has_text_frame else "")
                problemes.append("diapo %d : chevauche le pied de page (bas=%.2f) %r"
                                % (n, y + h, txt))

            if not sh.has_text_frame:
                continue
            tf = sh.text_frame
            if not tf.text.strip():
                continue
            dispo_h = h - (pouces(tf.margin_top) or 0) - (pouces(tf.margin_bottom) or 0)
            dispo_w = w - (pouces(tf.margin_left) or 0) - (pouces(tf.margin_right) or 0)
            if dispo_h <= 0.05:
                continue
            for p in tf.paragraphs:
                txt = "".join(r.text for r in p.runs)
                tailles = [r.font.size.pt for r in p.runs if r.font.size]
                if not txt.strip() or not tailles:
                    continue
                t = max(tailles)
                car = max(1.0, (dispo_w * 72.0) / (ratio_texte(t) * t))
                lignes = max(1, int(math.ceil(len(txt) / car)))
                besoin = lignes * t * INTER / 72.0
                if besoin > dispo_h + 0.04:
                    problemes.append("diapo %d : texte trop grand (%.2f\" pour %.2f\") %r"
                                    % (n, besoin, dispo_h, txt[:55]))

    nom = os.path.basename(chemin)
    if problemes:
        print("%s : %d probleme(s)" % (nom, len(problemes)))
        for p in problemes:
            print("  -", p)
        return 1
    print("%s : aucun probleme de mise en page sur %d diapositives."
          % (nom, len(prs.slides._sldIdLst)))
    return 0


if __name__ == "__main__":
    sys.exit(main())
