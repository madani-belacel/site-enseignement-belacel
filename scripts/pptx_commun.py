#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Primitives de mise en page PowerPoint, partagées par les générateurs
de présentation du site.

Pourquoi un module commun : le cours 01 de l'Intelligence Artificielle et le
cours 01 des TIC doivent avoir la même charte, les mêmes logos et le même
mécanisme d'ajustement des tailles. Dupliquer ce code dans chaque script
aurait fini par les faire diverger — c'est exactement le problème que le
générateur de l'IA avait avec le contenu du cours.

Contient :
  - la palette (couleurs du site) et les réglages de lisibilité ;
  - les primitives de formes, de texte et de schémas ;
  - la couverture commune avec les deux logos de l'institution ;
  - la calibration de la largeur d'un caractère, mesurée sur le rendu réel,
    qui permet de choisir une taille de police tenant dans une boîte.

Ne dépend que de python-pptx et de Pillow.
"""
import math
import os

from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.util import Inches, Pt

# ---------------------------------------------------------------------------
# Palette (couleurs du site)
# ---------------------------------------------------------------------------
NAVY = RGBColor(0x12, 0x2A, 0x4F)
NAVY_2 = RGBColor(0x1A, 0x38, 0x64)
BLUE = RGBColor(0x1F, 0x5C, 0x99)
BLUE_L = RGBColor(0xE4, 0xEF, 0xF9)
GREEN = RGBColor(0x1B, 0x7A, 0x52)
GREEN_L = RGBColor(0xE4, 0xF5, 0xEC)
AMBER = RGBColor(0xC2, 0x6A, 0x0A)
AMBER_L = RGBColor(0xFD, 0xF0, 0xDD)
GREY = RGBColor(0x5B, 0x68, 0x79)
GREY_L = RGBColor(0xF3, 0xF6, 0xFA)
LINE_C = RGBColor(0xD3, 0xDE, 0xEA)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
DARK = RGBColor(0x1B, 0x24, 0x33)
BLEU_CLAIR = RGBColor(0x9E, 0xBA, 0xDA)
TITRE_CLAIR = RGBColor(0xC5, 0xD6, 0xEA)
GRIS_PALE = RGBColor(0x7E, 0x93, 0xAD)
FILET_MARIN = RGBColor(0x2C, 0x45, 0x66)

FONT = "Calibri"

# style -> (fond, bordure, couleur du titre, couleur de la legende)
STYLES = {
    "primary": (BLUE_L, BLUE, NAVY, GREY),
    "alt": (GREEN_L, GREEN, RGBColor(0x0F, 0x4A, 0x32), GREY),
    "warn": (AMBER_L, AMBER, RGBColor(0x7A, 0x42, 0x05), GREY),
    "plain": (GREY_L, LINE_C, DARK, GREY),
    "ghost": (WHITE, GREY, GREY, GREY),
}

# ---------------------------------------------------------------------------
# Réglages de lisibilité
# ---------------------------------------------------------------------------
# Une salle de 30 places demande environ 20 pt pour un texte lu au tableau.
# Regroupés ici pour être ajustés sans toucher au reste.
TAILLE_TITRE = 30        # titre d'idee / de section
TAILLE_TEXTE_MAX = 19.0  # texte courant : plafond
TAILLE_PASTILLE = 12     # badge numerote
TAILLE_BOITE = 15.5      # texte principal des schemas
TAILLE_BOITE_SUB = 12.0  # legende des schemas
TAILLE_PIED = 9          # pied de page
TAILLE_ARROW = 21        # fleches entre boites

# La presentation du module IA est en francais seul (2eme annee PEP, filiere
# Francais). Mettre True pour revenir a une version bilingue.
BILINGUE = False

LARGEUR, HAUTEUR = 13.333, 7.5
CONTENU_G, CONTENU_D = 0.85, 11.63   # marge gauche, largeur utile du texte
LIGNE_PIED = 6.92                     # filet au-dessus du pied de page


# ---------------------------------------------------------------------------
# Calibration du texte
# ---------------------------------------------------------------------------
def ratio_texte(taille):
    """Largeur moyenne d'un caractere, rapportee a la taille de police.

    Calibre sur deux points mesures au rendu (LibreOffice -> PDF) :
      - 11 pt -> 0.65 (le texte debordait du cadre avec 0.50)
      - 38 pt -> 0.42 (le titre de couverture tient sur une seule ligne)
    Un ratio unique ne suffit pas : plus la police est grande, moins les
    espaces inter-caracteres fixes pesent en proportion.
    """
    return max(0.40, min(0.70, 1.094 - 0.427 * math.log10(max(1.0, taille))))


def lignes_requises(contenu, largeur_in, taille, ratio=None):
    r = ratio_texte(taille) if ratio is None else ratio
    car_par_ligne = max(1.0, (largeur_in * 72.0) / (r * taille))
    return max(1, int(math.ceil(len(contenu) / car_par_ligne)))


def hauteur_requise(contenu, largeur_in, taille, interligne=1.15, ratio=None):
    """Hauteur (en pouces) qu'occupe `contenu` a `taille`, dans `largeur_in`."""
    return (lignes_requises(contenu, largeur_in, taille, ratio)
            * taille * interligne / 72.0)


def taille_qui_tient(contenu, largeur_in, hauteur_in, maxi=None, mini=8.5,
                     interligne=1.15, marge=0.94):
    """Plus grande taille de police (<= maxi) telle que le texte tienne dans
    la boite, avec une marge de securite."""
    if maxi is None:
        maxi = TAILLE_TEXTE_MAX
    for pas in range(int(round((maxi - mini) * 2)) + 1):
        taille = round(maxi - pas / 2.0, 1)
        if hauteur_requise(contenu, largeur_in, taille, interligne) <= hauteur_in * marge:
            return taille
    return mini


# ---------------------------------------------------------------------------
# Primitives
# ---------------------------------------------------------------------------
def add_shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0,
              radius=None, dash=False):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
        if dash:
            try:
                from pptx.enum.dml import MSO_LINE_DASH_STYLE
                sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            except Exception:
                pass
    try:
        sh.shadow.inherit = False
    except Exception:
        pass
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    return sh


def add_textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def add_paragraph(tf, runs, size=13, bold=False, italic=False, color=DARK,
                  align=PP_ALIGN.LEFT, space_before=0, space_after=4,
                  line_spacing=1.12, font=FONT):
    if len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    if isinstance(runs, str):
        runs = [(runs, {})]
    for text, opt in runs:
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = opt.get("font", font)
        f.size = Pt(opt.get("size", size))
        f.bold = opt.get("bold", bold)
        f.italic = opt.get("italic", italic)
        f.color.rgb = opt.get("color", color)
    return p


def add_footer(slide, page_no, texte=""):
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.75, LIGNE_PIED, 11.83, 0.012, fill=LINE_C)
    tf = add_textbox(slide, 0.75, 7.02, 9.5, 0.3)
    add_paragraph(tf, [(texte, {"size": TAILLE_PIED, "color": GREY})], space_after=0)
    tf2 = add_textbox(slide, 11.0, 7.02, 1.58, 0.3)
    add_paragraph(tf2, [(str(page_no), {"size": TAILLE_PIED, "bold": True, "color": BLUE})],
                  align=PP_ALIGN.RIGHT, space_after=0)


# ---------------------------------------------------------------------------
# Schémas
# ---------------------------------------------------------------------------
def dbox(slide, x, y, w, h, titre, legende="", style="primary"):
    fill, line_c, txt_c, sub_c = STYLES.get(style, STYLES["primary"])
    sh = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, x, y, w, h,
                   fill=fill, line=line_c, line_w=1.25, radius=0.14,
                   dash=(style == "ghost"))
    tf = sh.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = MSO_ANCHOR.MIDDLE
    tf.margin_left = Inches(0.06)
    tf.margin_right = Inches(0.06)
    tf.margin_top = Inches(0.04)
    tf.margin_bottom = Inches(0.04)
    return sh, tf


def tailles_boite(box_w):
    """Les boites n'ont pas toutes la meme largeur : six boites sur une rangee
    ne peuvent pas prendre la meme taille de police que trois. plutot que de
    tout rapetisser, on adapte chaque rangee."""
    if box_w >= 2.30:
        return TAILLE_BOITE, TAILLE_BOITE_SUB
    if box_w >= 1.90:
        return TAILLE_BOITE - 1.5, TAILLE_BOITE_SUB - 1.0
    return TAILLE_BOITE - 3.0, TAILLE_BOITE_SUB - 2.0


def draw_flow(slide, elements, x, y, w, h):
    """elements : liste de ('box', main, sub, style) ou ('arrow', symbole)."""
    n_boxes = sum(1 for e in elements if e[0] == "box")
    n_conn = len(elements) - n_boxes
    conn_w, gap = (0.0, 0.18) if n_conn == 0 else (0.42, 0.0)
    box_w = (w - n_conn * conn_w - (n_boxes - 1) * gap) / n_boxes
    t_main, t_sub = tailles_boite(box_w)

    cx = x
    for i, e in enumerate(elements):
        if e[0] == "box":
            _, main, sub, style = e
            _, tf = dbox(slide, cx, y, box_w, h, main, sub, style)
            fill, line_c, txt_c, sub_c = STYLES.get(style, STYLES["primary"])
            add_paragraph(tf, [(main, {"size": t_main, "bold": True, "color": txt_c})],
                          align=PP_ALIGN.CENTER, space_after=0, line_spacing=1.0)
            if sub:
                add_paragraph(tf, [(sub, {"size": t_sub, "color": sub_c})],
                              align=PP_ALIGN.CENTER, space_before=3,
                              space_after=0, line_spacing=1.0)
            cx += box_w
            if i < len(elements) - 1:
                cx += gap
        else:
            _, sym = e
            cw = conn_w if conn_w else 0.42
            tf = add_textbox(slide, cx, y, cw, h, anchor=MSO_ANCHOR.MIDDLE)
            add_paragraph(tf, [(sym, {"size": TAILLE_ARROW, "bold": True, "color": BLUE})],
                          align=PP_ALIGN.CENTER, space_after=0, line_spacing=1.0)
            cx += conn_w


# ---------------------------------------------------------------------------
# Logos et couverture
# ---------------------------------------------------------------------------
def logo(slide, chemin, x, y, h, dpi=300, cache=None):
    """Pose un logo en conservant ses proportions ; renvoie sa largeur.

    python-pptx ne lit ni le .webp ni l'avif : le site les propose pourtant
    en version moderne. On prend donc la version PNG/JPEG, et on la retaille
    a la definition utile (300 ppp) si elle est beaucoup plus grande.
    """
    from PIL import Image
    try:
        with Image.open(chemin) as im:
            w_px, h_px = im.size
    except Exception as e:
        print("  logo illisible : %s (%s)" % (chemin, e))
        return 0.0
    if not h_px:
        return 0.0

    cible_h = max(64, int(round(h * dpi)))
    if cache and h_px > cible_h * 1.5:
        os.makedirs(cache, exist_ok=True)
        nom = os.path.splitext(os.path.basename(chemin))[0] + ".png"
        leger = os.path.join(cache, nom)
        if not os.path.exists(leger) or os.path.getmtime(leger) < os.path.getmtime(chemin):
            with Image.open(chemin) as im:
                im = im.convert("RGBA") if im.mode in ("RGBA", "LA", "P") else im.convert("RGB")
                im.thumbnail((cible_h * 4, cible_h), Image.LANCZOS)
                im.save(leger, optimize=True)
            print("  logo allégé :", nom)
        chemin = leger
        with Image.open(chemin) as im:
            w_px, h_px = im.size

    w = h * w_px / float(h_px)
    slide.shapes.add_picture(chemin, Inches(x), Inches(y), Inches(w), Inches(h))
    return w


def couverture(slide, logo_univ, logo_fle, titre, sous_titre, badge,
               auteur, mention, adresse=None, cache=None):
    """Page de garde commune : bandeau blanc avec les deux logos de
    l'institution, puis corps marine.

    Le logo de la faculte est un JPEG a fond blanc : le bandeau blanc est ce
    qui evite un rectangle gris visible sur le fond marine.
    """
    h_bande = 1.50
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, 0, LARGEUR, h_bande, fill=WHITE)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, h_bande, LARGEUR, 0.045, fill=AMBER)

    h_logo = 1.06
    y_logo = (h_bande - h_logo) / 2
    x = 0.80
    if logo_univ:
        x += logo(slide, logo_univ, x, y_logo, h_logo, cache=cache) + 0.34
    if logo_fle:
        x += logo(slide, logo_fle, x, y_logo, h_logo, cache=cache) + 0.46
    add_shape(slide, MSO_SHAPE.RECTANGLE, x, 0.48, 0.012, 0.54, fill=LINE_C)

    tf = add_textbox(slide, x + 0.30, 0.42, LARGEUR - x - 1.10, 0.68,
                     anchor=MSO_ANCHOR.MIDDLE)
    add_paragraph(tf, [("Université de Mostaganem  ·  Faculté des Langues Étrangères",
                        {"size": 12.5, "bold": True, "color": NAVY})], space_after=2)
    add_paragraph(tf, [(mention, {"size": 11, "color": GREY})], space_after=0)

    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, h_bande + 0.045, LARGEUR,
              HAUTEUR - h_bande - 0.045, fill=NAVY)
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0, h_bande + 0.045, 0.18,
              HAUTEUR - h_bande - 0.045, fill=AMBER)

    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.0, 2.42, 1.6, 0.07, fill=AMBER)

    tf = add_textbox(slide, 1.0, 2.72, 11.3, 1.1)
    add_paragraph(tf, [(titre, {"size": 40, "bold": True, "color": WHITE})],
                  line_spacing=1.0, space_after=0)

    if sous_titre:
        tf = add_textbox(slide, 1.0, 3.96, 10.9, 0.62)
        add_paragraph(tf, [(sous_titre, {"size": TAILLE_TEXTE_MAX - 2.0,
                                          "italic": True, "color": BLEU_CLAIR})],
                      space_after=0)

    if badge:
        add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, 1.0, 4.72, 5.4, 0.58,
                  fill=AMBER, radius=0.35)
        tf = add_textbox(slide, 1.0, 4.72, 5.4, 0.58, anchor=MSO_ANCHOR.MIDDLE)
        add_paragraph(tf, [(badge, {"size": 15, "bold": True, "color": NAVY})],
                      align=PP_ALIGN.CENTER, space_after=0)

    add_shape(slide, MSO_SHAPE.RECTANGLE, 1.0, 5.72, 11.33, 0.012, fill=FILET_MARIN)

    tf = add_textbox(slide, 1.0, 5.96, 11.3, 1.0)
    add_paragraph(tf, [(auteur, {"size": 16, "bold": True, "color": WHITE})],
                  space_after=3)
    if adresse:
        add_paragraph(tf, [(adresse, {"size": 10, "color": GRIS_PALE})], space_after=0)
    return


# ---------------------------------------------------------------------------
# Aides pour les générateurs
# ---------------------------------------------------------------------------
def plain(fragment):
    """Balisage HTML -> texte brut, balises de mise en forme comprises."""
    import re
    from html import unescape
    t = re.sub(r"<br\s*/?>", " ", fragment)
    t = re.sub(r"<[^>]+>", "", t)
    return re.sub(r"\s+", " ", unescape(t)).strip()


def nouveau_conteneur():
    from pptx import Presentation
    prs = Presentation()
    prs.slide_width = Inches(LARGEUR)
    prs.slide_height = Inches(HAUTEUR)
    return prs, prs.slide_layouts[6]


def ecrire_notes(slide, texte):
    slide.notes_slide.notes_text_frame.text = texte.strip()
