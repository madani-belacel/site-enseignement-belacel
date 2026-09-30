#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Génère QR_IA_01_Introduction_IA.pptx
  - Diapo 1     : page de garde
  - Diapos 2-11 : une idée par diapositive (10 idées)
  - Notes       : les Questions / Réponses (FR + EN) — 32 paires

11 diapositives (1 couverture + 10 idées), format 16:9 (13,333 x 7,5 po).
Sortie : Presentations_PPTX/QR_IA_01_Introduction_IA.pptx

Les titres, les explications et les Q/R sont RELUS dans
QR_IA_01_Introduction_IA.html : le PowerPoint ne peut donc pas diverger du
cours. Seuls les schémas (IDEES[...]["flows"]) sont dessinés ici, parce
qu'ils sont simplifiés pour la projection. Si le HTML est illisible, le
script retombe sur les textes inscribed plus bas dans le fichier.
"""

import math
import os
import re
from html import unescape

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE

try:
    from pptx.enum.dml import MSO_LINE_DASH_STYLE
    HAS_DASH = True
except Exception:
    HAS_DASH = False

# ----------------------------------------------------------------------
# Palette
# ----------------------------------------------------------------------
NAVY    = RGBColor(0x12, 0x2A, 0x4F)
NAVY_2  = RGBColor(0x1A, 0x38, 0x64)
BLUE    = RGBColor(0x1F, 0x5C, 0x99)
BLUE_L  = RGBColor(0xE4, 0xEF, 0xF9)
GREEN   = RGBColor(0x1B, 0x7A, 0x52)
GREEN_L = RGBColor(0xE4, 0xF5, 0xEC)
AMBER   = RGBColor(0xC2, 0x6A, 0x0A)
AMBER_L = RGBColor(0xFD, 0xF0, 0xDD)
GREY    = RGBColor(0x5B, 0x68, 0x79)
GREY_L  = RGBColor(0xF3, 0xF6, 0xFA)
LINE_C  = RGBColor(0xD3, 0xDE, 0xEA)
WHITE   = RGBColor(0xFF, 0xFF, 0xFF)
DARK    = RGBColor(0x1B, 0x24, 0x33)

FONT = "Calibri"

STYLES = {
    "primary": (BLUE_L,  BLUE,  NAVY,                    GREY),
    "alt":     (GREEN_L, GREEN, RGBColor(0x0F, 0x4A, 0x32), GREY),
    "warn":    (AMBER_L, AMBER, RGBColor(0x7A, 0x42, 0x05), GREY),
    "plain":   (GREY_L,  LINE_C, DARK,                   GREY),
    "ghost":   (WHITE,   GREY,  GREY,                    GREY),
}

# ----------------------------------------------------------------------
# Helpers
# ----------------------------------------------------------------------
def add_shape(slide, kind, x, y, w, h, fill=None, line=None, line_w=1.0,
              radius=None, dash=False):
    sh = slide.shapes.add_shape(kind, Inches(x), Inches(y), Inches(w), Inches(h))
    if fill is None:
        sh.fill.background()
    else:
        sh.fill.solid()
        sh.fill.fore_color.rgb = fill
    if line is None:
        sh.line.fill.background()
    else:
        sh.line.color.rgb = line
        sh.line.width = Pt(line_w)
        if dash and HAS_DASH:
            try:
                sh.line.dash_style = MSO_LINE_DASH_STYLE.DASH
            except Exception:
                pass
    try:
        sh.shadow.inherit = False
    except Exception:
        pass
    if radius is not None and kind == MSO_SHAPE.ROUNDED_RECTANGLE:
        try:
            sh.adjustments[0] = radius
        except Exception:
            pass
    return sh


def add_textbox(slide, x, y, w, h, anchor=MSO_ANCHOR.TOP):
    tb = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    tf.vertical_anchor = anchor
    return tf


def add_paragraph(tf, runs, size=13, bold=False, italic=False, color=DARK,
                  align=PP_ALIGN.LEFT, space_before=0, space_after=4,
                  line_spacing=1.12, font=FONT):
    if len(tf.paragraphs) == 1 and not tf.paragraphs[0].runs:
        p = tf.paragraphs[0]
    else:
        p = tf.add_paragraph()
    p.alignment = align
    p.space_before = Pt(space_before)
    p.space_after = Pt(space_after)
    p.line_spacing = line_spacing
    if isinstance(runs, str):
        runs = [(runs, {})]
    for text, opt in runs:
        r = p.add_run()
        r.text = text
        f = r.font
        f.name = opt.get("font", font)
        f.size = Pt(opt.get("size", size))
        f.bold = opt.get("bold", bold)
        f.italic = opt.get("italic", italic)
        f.color.rgb = opt.get("color", color)
    return p


def ratio_texte(taille):
    """Largeur moyenne d'un caractere, rapportee a la taille de police.

    Calibre sur deux points mesures au rendu (LibreOffice -> PDF) :
      - 11 pt  -> 0.65 (l'idee 3 debordait avec 0.50)
      - 38 pt  -> 0.42 (le titre de couverture tient sur une ligne)
    Un ratio unique ne peut pas servir : plus la police est grande, plus les
    espaces inter-caracteres fixes pesent peu en proportion.
    """
    return max(0.40, min(0.70, 1.094 - 0.427 * math.log10(max(1.0, taille))))


def hauteur_requise(contenu, largeur_in, taille, interligne=1.15):
    """Hauteur (en pouces) qu'occupe `contenu` a `taille`, dans `largeur_in`."""
    car_par_ligne = max(1.0, (largeur_in * 72.0) / (ratio_texte(taille) * taille))
    lignes = max(1, int(math.ceil(len(contenu) / car_par_ligne)))
    return lignes * taille * interligne / 72.0


def taille_qui_tient(contenu, largeur_in, hauteur_in, maxi=13.5, mini=8.5,
                     interligne=1.15, marge=0.94):
    """Plus grande taille de police (<= maxi) telle que le texte tienne dans
    la boite, avec une marge de securite."""
    for pas in range(int(round((maxi - mini) * 2)) + 1):
        taille = round(maxi - pas / 2.0, 1)
        if hauteur_requise(contenu, largeur_in, taille, interligne) <= hauteur_in * marge:
            return taille
    return mini


def draw_flow(slide, elements, x, y, w, h):
    """elements : liste de ('box', main, sub, style) ou ('arrow', symbole)."""
    n_boxes = sum(1 for e in elements if e[0] == "box")
    n_conn = len(elements) - n_boxes
    if n_conn == 0:
        conn_w, gap = 0.0, 0.18
    else:
        conn_w, gap = 0.42, 0.0
    box_w = (w - n_conn * conn_w - (n_boxes - 1) * gap) / n_boxes

    # Les boites n'ont pas toutes la meme largeur : sur l'idee 4 elles sont
    # six sur une rangee. On grossit donc le texte dans les boites larges et
    # on le reduit dans les boites etroites, plutot que de tout rapetisser.
    if box_w >= 2.30:
        t_main, t_sub = TAILLE_BOITE, TAILLE_BOITE_SUB
    elif box_w >= 1.90:
        t_main, t_sub = TAILLE_BOITE - 1.5, TAILLE_BOITE_SUB - 1.0
    else:
        t_main, t_sub = TAILLE_BOITE - 3.0, TAILLE_BOITE_SUB - 2.0

    cx = x
    for i, e in enumerate(elements):
        if e[0] == "box":
            _, main, sub, style = e
            fill, line_c, txt_c, sub_c = STYLES[style]
            dashed = (style == "ghost")
            sh = add_shape(slide, MSO_SHAPE.ROUNDED_RECTANGLE, cx, y, box_w, h,
                           fill=fill, line=line_c, line_w=1.25,
                           radius=0.14, dash=dashed)
            tf = sh.text_frame
            tf.word_wrap = True
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = Inches(0.06)
            tf.margin_right = Inches(0.06)
            tf.margin_top = Inches(0.04)
            tf.margin_bottom = Inches(0.04)
            add_paragraph(tf, [(main, {"size": t_main, "bold": True, "color": txt_c})],
                          align=PP_ALIGN.CENTER, space_after=0, line_spacing=1.0)
            if sub:
                add_paragraph(tf, [(sub, {"size": t_sub, "color": sub_c})],
                              align=PP_ALIGN.CENTER, space_before=3,
                              space_after=0, line_spacing=1.0)
            cx += box_w
            if i < len(elements) - 1:
                cx += gap
        else:
            _, sym = e
            cw = conn_w if conn_w else 0.42
            tb = slide.shapes.add_textbox(Inches(cx), Inches(y), Inches(cw), Inches(h))
            tf = tb.text_frame
            tf.vertical_anchor = MSO_ANCHOR.MIDDLE
            tf.margin_left = 0; tf.margin_right = 0
            tf.margin_top = 0;  tf.margin_bottom = 0
            add_paragraph(tf, [(sym, {"size": 21, "bold": True, "color": BLUE})],
                          align=PP_ALIGN.CENTER, space_after=0, line_spacing=1.0)
            cx += conn_w


def add_footer(slide, page_no):
    add_shape(slide, MSO_SHAPE.RECTANGLE, 0.75, 6.92, 11.83, 0.012, fill=LINE_C)
    tf = add_textbox(slide, 0.75, 7.02, 9.5, 0.3)
    add_paragraph(tf, [("Q/R — Cours 01 · Introduction à l'IA   |   Dr. Madani BELACEL — Université de Mostaganem",
                        {"size": 9, "color": GREY})], space_after=0)
    tf2 = add_textbox(slide, 11.0, 7.02, 1.58, 0.3)
    add_paragraph(tf2, [(str(page_no), {"size": 9, "bold": True, "color": BLUE})],
                  align=PP_ALIGN.RIGHT, space_after=0)


# ----------------------------------------------------------------------
# 1. Lecture du cours HTML (source de vérité)
# ----------------------------------------------------------------------
DOSSIER = os.path.dirname(os.path.abspath(__file__))
COURS_HTML = os.path.join(DOSSIER, "QR_IA_01_Introduction_IA.html")
# python-pptx n'accepte que PNG / JPEG / GIF / BMP / TIFF : les versions .webp
# et .avif du site ne sont pas lisibles, on prend donc le PNG de l'université
# et le JPEG de la faculté.
IMAGES = os.path.abspath(os.path.join(DOSSIER, "..", "..", "images"))
LOGO_UNIV = os.path.join(IMAGES, "Université_de_Mostaganem.png")
LOGO_FLE = os.path.join(IMAGES, "LOGO-FLE-UNIV-Mosta.jpeg")
DOSSIER_LOGOS = os.path.join(DOSSIER, "logos_pptx")


def _propre(fragment):
    """Balisage HTML -> texte brut, balises de mise en forme comprises."""
    t = re.sub(r"<br\s*/?>", " ", fragment)
    t = re.sub(r"<[^>]+>", "", t)
    return re.sub(r"\s+", " ", unescape(t)).strip()


def lire_cours_html(chemin):
    """Relit les 10 idées du cours : titre FR/EN, explication FR/EN, Q/R FR/EN.

    Le repérage se fait sur les <article class="qr-idea"> ; un article ne
    contenant jamais un autre article, le premier </article> encountered le
    ferme toujours. Un <div class="qr-item"> ne contient pas plus de <p>,
    donc le premier </div> le ferme également.
    """
    s = open(chemin, encoding="utf-8").read()
    idees = []
    for art in re.findall(r'<article class="qr-idea">.*?</article>', s, re.S):
        def un(regex, defaut=""):
            m = re.search(regex, art, re.S)
            return _propre(m.group(1)) if m else defaut

        idees.append({
            "fr_t": un(r'<span class="lang-fr">(.*?)</span>\s*</h3>'),
            "en_t": un(r'<h3>.*?<span class="lang-en">(.*?)</span>'),
            "fr_x": un(r'<p class="qr-expl lang-fr">(.*?)</p>'),
            "en_x": un(r'<p class="qr-expl lang-en">(.*?)</p>'),
            "qa_fr": [], "qa_en": [],
        })
        for item in re.findall(r'<div class="qr-item">(.*?)</div>', art, re.S):
            # une <div class="qr-item"> peut contenir plusieurs paires (cas de
            # l'idee 3, ou Q3/R3 et Q4/R4 partagent la meme boite) : on
            # extrait donc toutes les paires, pas seulement la premiere.
            for langue, cle in (("fr", "qa_fr"), ("en", "qa_en")):
                qs = [_propre(x) for x in
                      re.findall(r'<p class="qr-q lang-%s">(.*?)</p>' % langue, item, re.S)]
                rs = [_propre(x) for x in
                      re.findall(r'<p class="qr-r lang-%s">(.*?)</p>' % langue, item, re.S)]
                for q, r in zip(qs, rs):
                    idees[-1][cle].append((q, r))
    return idees


# ----------------------------------------------------------------------
# 2. Le contenu
# ----------------------------------------------------------------------
B = lambda main, sub, style="primary": ("box", main, sub, style)
A = lambda sym: ("arrow", sym)
IDEES = [
    dict(
        fr_t="Qu'est-ce que l'Intelligence Artificielle ?",
        en_t="What is Artificial Intelligence?",
        fr_x="L'IA est un programme informatique qui peut apprendre, reconnaître des choses et prendre des décisions, un peu comme un humain. Ce n'est pas de la magie : c'est basé sur des données et des règles.",
        en_x="AI is a computer program that can learn, recognize things and make decisions, a bit like a human would. It is not magic: it is based on data and rules.",
        flows=[[B("Exemples", "Données", "primary"), A("→"),
                B("Régularités", "Les règles qu'il trouve", "alt"), A("→"),
                B("Une réponse", "Sur un cas nouveau", "warn")]],
        qa_fr=[("Q1. Qu'est-ce que l'IA en une phrase ?",
                "R1. L'IA est un programme informatique qui peut apprendre à partir de données et prendre des décisions tout seul."),
               ("Q2. L'IA est-elle un robot ?",
                "R2. Pas toujours. L'IA est un logiciel. Un robot est une machine qui peut bouger, et un robot peut contenir de l'IA."),
               ("Q3. L'IA pense-t-elle comme un humain ?",
                "R3. Non. L'IA imite certaines tâches humaines, comme reconnaître ou décider, mais elle ne ressent pas et ne comprend pas comme un humain.")],
        qa_en=[("Q1. What is AI in one sentence?",
                "A1. AI is a computer program that can learn from data and make decisions by itself."),
               ("Q2. Is AI a robot?",
                "A2. Not always. AI is software. A robot is a machine that can move, and a robot can contain AI."),
               ("Q3. Does AI think like a human?",
                "A3. No. AI imitates some human tasks, like recognizing or deciding, but it does not feel or understand the way a human does.")],
    ),
    dict(
        fr_t="Comment l'IA travaille-t-elle, en mots simples ?",
        en_t="How does AI work, in simple words?",
        fr_x="Un programme d'IA reçoit un très grand nombre d'exemples, chacun déjà étiqueté avec la bonne réponse. Il cherche la règle qui les relie, puis il applique cette règle à des cas qu'il n'a jamais vus. Personne n'écrit toutes les règles à la main : le programme les retrouve lui-même. Et il ne comprend pas vraiment : il reconnaît des formes et des statistiques, pas du sens.",
        en_x="An AI program is given a very large number of examples, each one already labelled with the right answer. It looks for the rule that links them, then applies that rule to cases it has never seen. Nobody writes all the rules by hand: the program finds them by itself. And it does not really understand — it recognises shapes and statistics, not meaning.",
        flows=[[B("1 000 photos de chats", "étiquetées « chat »", "primary"), A("+"),
                B("1 000 photos de chiens", "étiquetées « chien »", "primary"), A("→"),
                B("Une photo nouvelle", "jamais vue", "alt"), A("→"),
                B("chat ou chien", "la réponse", "warn")]],
        qa_fr=[("Q1. Est-ce que l'IA ressemble à Excel ?",
                "R1. Bien plus qu'on ne le croit. Dans Excel, 4 puis 8, on tire la petite poignée en bas à droite et Excel calcule 12 tout seul. L'IA fait la même chose avec des exemples : deux cas, elle trouve la règle qui les relie, puis elle l'applique au suivant."),
               ("Q2. Donnez-moi un exemple concret.",
                "R2. Un dossier de 1 000 photos de chats et 1 000 photos de chiens déjà étiquetées. On montre ensuite une photo jamais vue : le programme répond chat ou chien. Même idée pour un dossier de factures, de CV ou de sujets d'examen."),
               ("Q3. Que se passe-t-il si les exemples sont faux ?",
                "R3. L'IA recopie l'erreur. Des exemples de mauvaise qualité donnent des réponses de mauvaise qualité : la qualité des données compte plus que leur quantité.")],
        qa_en=[("Q1. Is AI like Excel?",
                "A1. Much more than you might think. In Excel you put 4 then 8, drag the small handle at the bottom-right corner, and Excel works out 12 by itself. An AI does the same with examples: two cases, it finds the rule, then applies it to the next one."),
               ("Q2. Give me a concrete example.",
                "A2. A folder with 1 000 cat photos and 1 000 dog photos, already labelled. Show it a photo it has never seen: it answers cat or dog. The same idea works for invoices, CVs or exam papers."),
               ("Q3. What happens if the examples are wrong?",
                "A3. The AI copies the mistake. Poor examples give poor answers, which is why the quality of the data matters more than the quantity.")],
    ),
    dict(
        fr_t="Deux familles d'IA : lui parler, ou coder avec elle",
        en_t="Two families of AI: talking to it, or coding with it",
        fr_x="D'un côté, les assistants de conversation — ChatGPT, Grok, DeepSeek — où vous écrivez une question en langue normale et vous obtenez une réponse rédigée, sans que rien ne soit touché sur votre ordinateur. De l'autre côté, les assistants de code — OpenCode, Cursor, GitHub Copilot — où l'IA travaille directement sur vos fichiers : elle peut les lire, en créer de nouveaux, les modifier, et analyser tout un dossier pour vous.",
        en_x="On one side, the conversational assistants — ChatGPT, Grok, DeepSeek — where you write a question in ordinary language and get a written answer, without anything on your computer being touched. On the other side, the coding assistants — OpenCode, Cursor, GitHub Copilot — where the AI works directly on your files: it can read them, create new ones, modify them, and analyse a whole folder for you.",
        flows=[[B("Assistants de conversation", "ChatGPT · Grok · DeepSeek\nVous écrivez → il répond en texte", "primary"),
                A("|"),
                B("Assistants de code", "OpenCode · Cursor · Copilot\nIl lit, crée, modifie vos fichiers", "warn")]],
        qa_fr=[("Q1. Quelle est la différence entre ChatGPT et Cursor ?",
                "R1. ChatGPT parle : il répond avec des phrases, dans la fenêtre de discussion. Cursor agit : il écrit et modifie directement les lignes de votre programme dans votre éditeur."),
               ("Q2. OpenCode, Cursor ou GitHub Copilot peuvent-ils toucher à mes fichiers ?",
                "R2. Oui, et il faut le savoir. Ils travaillent sur votre disque dur : ils lisent, créent, modifient et analysent un projet entier. C'est pourquoi on commence toujours dans un dossier de test."),
               ("Q3. OpenCode exige-t-il de savoir déjà programmer ?",
                "R3. Non, il vous aide à apprendre. Vous décrivez en mots ce que vous voulez, il propose le code et il explique chaque ligne."),
               ("Q4. Cursor repose-t-il sur le même principe que ChatGPT ?",
                "R4. Oui, exactement le même principe : des exemples, une règle, puis une application sur un cas nouveau. Ce qui change, ce n'est pas le principe, c'est ce qu'il a sous les yeux. ChatGPT répond à partir de ce qu'il a appris ; Cursor ouvre en plus vos fichiers, les lit, puis applique le même principe. C'est le même cerveau : on lui montre simplement un dossier en plus."),
               ("Q5. J'ai vu Cursor et OpenCode exécuter des scripts. C'est l'IA qui exécute ?",
                "R5. Non : c'est l'outil qui exécute, pas l'IA. L'IA écrit la commande, l'ordinateur la lance, et le résultat revient à l'IA. Cela forme une boucle — écrire, exécuter, corriger — que ChatGPT ne peut pas faire.")],
        qa_en=[("Q1. What is the difference between ChatGPT and Cursor?",
                "A1. ChatGPT talks: it answers with sentences inside the chat window. Cursor acts: it writes and modifies the lines of your program directly in your editor."),
               ("Q2. Can OpenCode, Cursor or GitHub Copilot touch my files?",
                "A2. Yes, and you must know it. They work on your hard disk: they read, create, modify and analyse a whole project. This is why you always start in a test folder."),
               ("Q3. Does OpenCode require me to already know programming?",
                "A3. No, it helps you learn. You describe in words what you want, it proposes the code and explains every line."),
               ("Q4. Does Cursor rest on the same principle as ChatGPT?",
                "A4. Yes, exactly the same principle: examples, a rule, then an application to a new case. What changes is not the principle, it is what it has in front of its eyes. ChatGPT answers from what it learned; Cursor also opens your files, reads them, then applies the same principle. Same brain — you simply hand it one more folder."),
               ("Q5. I have seen Cursor and OpenCode run scripts. Does the AI itself run them?",
                "A5. No: it is the tool that runs things, not the AI. The AI writes the command, the computer runs it, and the result comes back. That forms a loop — write, run, correct — which ChatGPT cannot do.")],
    ),
    dict(
        fr_t="Où rencontre-t-on l'IA ?",
        en_t="Where do we meet AI?",
        fr_x="On rencontre l'IA partout : dans ton téléphone (photos, voix), dans ton courriel (anti-spam), sur les réseaux sociaux (recommandations), dans les cartes en ligne, dans la météo et même dans le diagnostic médical.",
        en_x="AI is everywhere: in your phone (photos, voice), in your e-mail (spam filter), on social media (recommendations), in online maps, in the weather and even in medical diagnosis.",
        flows=[[B("Téléphone", "", "primary"), B("Courriel", "", "primary"),
                B("Cartes", "", "alt"), B("Météo", "", "alt"),
                B("Streaming", "", "warn"), B("Médecine", "", "warn")]],
        qa_fr=[("Q1. Où est l'IA dans un smartphone ?",
                "R1. Dans le déverrouillage par visage, le tri des photos, les assistants vocaux et le clavier prédictif."),
               ("Q2. L'IA existe-t-elle dans le courriel ?",
                "R2. Oui, le filtre anti-spam qui envoie le courrier indésirable dans le dossier spam est une petite IA."),
               ("Q3. L'IA est-elle réservée aux scientifiques ?",
                "R3. Non, tout le monde utilise l'IA tous les jours, souvent sans le savoir.")],
        qa_en=[("Q1. Where is AI in a smartphone?",
                "A1. In face unlock, photo sorting, voice assistants and the predictive keyboard."),
               ("Q2. Does AI exist in e-mail?",
                "A2. Yes, the spam filter that sends junk mail to the spam folder is a small AI."),
               ("Q3. Is AI only for scientists?",
                "A3. No, everyone uses AI every day, often without knowing it.")],
    ),
    dict(
        fr_t="Pourquoi l'IA est-elle possible aujourd'hui ?",
        en_t="Why is AI possible today?",
        fr_x="L'IA est devenue possible grâce à trois choses : beaucoup de données (photos, textes, vidéos), des ordinateurs très puissants et de meilleurs algorithmes.",
        en_x="AI became possible thanks to three things: a lot of data (photos, texts, videos), very powerful computers, and better algorithms.",
        flows=[[B("Plus de données", "sur Internet", "primary"), A("+"),
                B("Plus de puissance", "GPU, serveurs abordables", "primary"), A("+"),
                B("De meilleurs outils", "gratuits, dans le navigateur", "primary"), A("→"),
                B("L'IA pour tous", "", "warn")]],
        qa_fr=[("Q1. Quel rôle jouent les données ?",
                "R1. Les données sont le carburant de l'IA : plus la machine a de bons exemples, mieux elle apprend."),
               ("Q2. Quel rôle joue l'ordinateur ?",
                "R2. Les ordinateurs modernes sont assez rapides et puissants pour faire tous les calculs dont l'IA a besoin."),
               ("Q3. Les algorithmes ont-ils toujours existé ?",
                "R3. Beaucoup d'idées sont anciennes, mais aujourd'hui nous avons enfin les données et la puissance de calcul pour bien les utiliser.")],
        qa_en=[("Q1. What role does data play?",
                "A1. Data is the fuel of AI: the more good examples the AI has, the better it learns."),
               ("Q2. What role does the computer play?",
                "A2. Modern computers are fast and powerful enough to do all the calculations that AI needs."),
               ("Q3. Were the algorithms always here?",
                "A3. Many ideas are old, but today we finally have the data and the computing power to use them well.")],
    ),
    dict(
        fr_t="Le rôle des données",
        en_t="The role of data",
        fr_x="L'IA a besoin d'exemples pour apprendre. Si on montre 1 000 photos de chats et 1 000 photos de chiens, l'IA apprend à les distinguer. Pas de données, pas d'IA. Et la qualité compte plus que la quantité.",
        en_x="AI needs examples to learn. If you show 1 000 photos of cats and 1 000 photos of dogs, the AI learns to tell them apart. No data, no AI. And quality matters more than quantity.",
        flows=[
            [B("Peu mais propres", "500 bons exemples", "alt"), A("→"),
             B("Réponse fiable", "", "alt")],
            [B("Beaucoup mais sales", "10 000, la moitié mal étiquetée", "warn"), A("→"),
             B("Réponse confuse", "", "ghost")],
        ],
        qa_fr=[("Q1. Qu'est-ce que les données d'entraînement ?",
                "R1. Ce sont les exemples utilisés pour apprendre à l'IA, comme des photos de chiens et de chats avec leurs étiquettes."),
               ("Q2. Pourquoi l'IA a-t-elle besoin de beaucoup de données ?",
                "R2. Plus d'exemples de qualité aident l'IA à mieux trouver les régularités et à faire moins d'erreurs."),
               ("Q3. L'IA peut-elle apprendre sans données ?",
                "R3. Non. Sans données, il n'y a rien à partir de quoi apprendre.")],
        qa_en=[("Q1. What is training data?",
                "A1. The examples used to teach the AI, like photos of dogs and cats with their labels."),
               ("Q2. Why does AI need a lot of data?",
                "A2. More good examples help the AI find better patterns and make fewer mistakes."),
               ("Q3. Can AI learn without data?",
                "A3. No. Without data, there is nothing to learn from.")],
    ),
    dict(
        fr_t="L'IA dans l'éducation",
        en_t="AI in education",
        fr_x="L'IA aide les étudiants : elle traduit, corrige la grammaire, explique les exercices, crée des exercices et aide les enseignants à gagner du temps. Mais elle ne remplace ni l'effort ni le professeur.",
        en_x="AI helps students: it translates, corrects grammar, explains exercises, creates exercises, and helps teachers save time. But it replaces neither effort nor the teacher.",
        flows=[[B("Préparer", "résumé, plan", "primary"), A("→"),
                B("S'entraîner", "QCM, fiches", "alt"), A("→"),
                B("Rédiger", "corriger, traduire", "primary"), A("→"),
                B("Toujours vérifier", "", "warn")]],
        qa_fr=[("Q1. Comment l'IA peut-elle t'aider à apprendre l'anglais ?",
                "R1. Elle peut traduire, corriger tes phrases et te donner des exemples de dialogues."),
               ("Q2. L'IA peut-elle faire tes devoirs ?",
                "R2. Elle peut aider, mais tu dois apprendre. Si elle pense à ta place, tu n'apprendras rien."),
               ("Q3. L'IA peut-elle remplacer un enseignant ?",
                "R3. Non. C'est un outil qui aide, mais seul un humain peut comprendre, motiver et guider un étudiant.")],
        qa_en=[("Q1. How can AI help you learn English?",
                "A1. It can translate, correct your sentences and give you examples of dialogues."),
               ("Q2. Can AI write your homework?",
                "A2. It can help, but you must learn. If it thinks for you, you will not learn anything."),
               ("Q3. Can AI replace a teacher?",
                "A3. No. It is a tool that helps, but only a human can understand, motivate and guide a student.")],
    ),
    dict(
        fr_t="Est-ce que l'IA comprend vraiment ?",
        en_t="Does AI really understand?",
        fr_x="L'IA réagit aux régularités dans les données. Elle peut écrire une phrase correcte, mais elle ne comprend pas vraiment comme un humain. Elle n'a ni sentiments ni convictions.",
        en_x="AI reacts to patterns in data. It can write a correct sentence, but it does not really understand like a human. It has no feelings and no beliefs.",
        flows=[[B("Texte", "", "primary"), A("→"),
                B("Formes et statistiques", "ce que l'IA calcule", "alt"), A("≠"),
                B("Du sens", "ce que l'IA n'atteint pas", "ghost")]],
        qa_fr=[("Q1. L'IA comprend-elle les mots ?",
                "R1. Elle voit des régularités et des statistiques, pas le vrai sens. Elle prédit quel mot vient probablement ensuite."),
               ("Q2. L'IA a-t-elle des sentiments ?",
                "R2. Non. Elle peut imiter des émotions dans un texte, mais elle ne ressent rien."),
               ("Q3. Pourquoi l'IA fait-elle parfois des erreurs ?",
                "R3. Parce qu'elle ne connaît que ce qu'elle a vu. Si la situation est nouvelle ou si les données contiennent des erreurs, elle peut se tromper.")],
        qa_en=[("Q1. Does AI understand words?",
                "A1. It sees patterns and statistics, not real meaning. It predicts which word probably comes next."),
               ("Q2. Does AI have feelings?",
                "A2. No. It can imitate emotions in a text, but it does not feel anything."),
               ("Q3. Why does AI sometimes make mistakes?",
                "A3. Because it only knows what it has seen. If the situation is new or the data has errors, it can be wrong.")],
    ),
    dict(
        fr_t="L'IA et l'éthique",
        en_t="AI and ethics",
        fr_x="L'IA peut être utile mais aussi dangereuse : elle peut être biaisée, se tromper ou être utilisée à mauvais escient. Un humain doit toujours rester responsable.",
        en_x="AI can be useful but also dangerous: it can be biased, wrong, or used badly. A human must always stay responsible.",
        flows=[[B("Est-ce vrai ?", "vérifier l'information", "warn"),
                B("Ai-je cité ?", "donner la source", "warn"),
                B("Est-ce le mien ?", "produire son propre travail", "warn")]],
        qa_fr=[("Q1. Qu'est-ce que le biais d'une IA ?",
                "R1. C'est quand l'IA a appris à partir d'exemples déséquilibrés et fait donc des choix injustes ou faux."),
               ("Q2. Qui est responsable d'une décision prise par l'IA ?",
                "R2. Les humains qui ont conçu, choisi et utilisé l'IA."),
               ("Q3. Doit-on faire entièrement confiance à l'IA ?",
                "R3. Non. On doit vérifier ses réponses, comme on vérifie une calculatrice ou un dictionnaire.")],
        qa_en=[("Q1. What is AI bias?",
                "A1. When the AI has learned from unbalanced examples, so it makes unfair or wrong choices."),
               ("Q2. Who is responsible for an AI decision?",
                "A2. The humans who designed, chose and used the AI."),
               ("Q3. Should we trust AI completely?",
                "A3. No. We should check its answers, just like we check a calculator or a dictionary.")],
    ),
    dict(
        fr_t="Une brève histoire de l'IA",
        en_t="A short history of AI",
        fr_x="Le terme « Intelligence Artificielle » est né en 1956. Après des hauts et des bas, l'IA est devenue énorme après 2010 grâce au Big Data et aux ordinateurs puissants. ChatGPT a montré l'IA à tout le monde en 2022.",
        en_x="The term 'Artificial Intelligence' was born in 1956. After ups and downs, AI became huge after 2010 thanks to Big Data and powerful computers. ChatGPT showed AI to everyone in 2022.",
        flows=[[B("1950", "Test de Turing", "plain"), A("→"),
                B("1997", "Deep Blue", "plain"), A("→"),
                B("2011", "Siri", "plain"), A("→"),
                B("2022…", "ChatGPT, Copilot", "warn")]],
        qa_fr=[("Q1. Quand le terme IA a-t-il été créé ?",
                "R1. En 1956, lors d'une conférence aux États-Unis."),
               ("Q2. Pourquoi ChatGPT est-il célèbre ?",
                "R2. Parce qu'il parle si naturellement que n'importe qui peut utiliser l'IA sans compétence particulière."),
               ("Q3. L'IA est-elle nouvelle ?",
                "R3. L'idée est ancienne, mais l'explosion réelle est très récente, depuis une quinzaine d'années environ.")],
        qa_en=[("Q1. When was the term AI created?",
                "A1. In 1956, at a conference in the United States."),
               ("Q2. Why is ChatGPT famous?",
                "A2. Because it talks so naturally that anyone can use AI without special skills."),
               ("Q3. Is AI new?",
                "A3. The idea is old, but the real explosion is very recent, in the last fifteen years or so.")],
    ),
]


# ----------------------------------------------------------------------
# Diapositives
# ----------------------------------------------------------------------

# Langues affichees. Le module est destine a la 2eme annee PEP (ENS, filière
# Francais) : la seance se donne en francais, la presentation n'affiche donc
# que le francais. Mettre ("fr", "en") pour revenir a la version bilingue.
LANGUES = ("fr",)
BILINGUE = len(LANGUES) > 1

# --- Reglages de lisibilite -------------------------------------------------
# Une salle de 30 places demande environ 20 pt pour un texte lu au tableau.
# Ces valeurs sont regroupees ici : on les ajuste sans toucher au reste.
TAILLE_TITRE = 30        # titre de l'idee
TAILLE_TEXTE_MAX = 19.0  # explication francaise : plafond
TAILLE_PASTILLE = 12     # badge « IDEE n »
TAILLE_BOITE = 15.5      # texte principal des schemas
TAILLE_BOITE_SUB = 12.0  # legende des schemas


def _logo(slide, chemin, x, y, h, dpi=300):
    """Place un logo en conservant ses proportions ; renvoie sa largeur.

    Le PNG de l'universite fait 960 x 931 px pour 1,06 po d'affichage, soit
    900 ppp : inutile et il alourdissait le .pptx de 250 Ko. On genere donc une
    version a la bonne definition dans logos_pptx/ (une seule fois).
    """
    from PIL import Image as _Image
    try:
        with _Image.open(chemin) as im:
            w_px, h_px = im.size
    except Exception:
        print("  logo introuvable :", chemin)
        return 0.0
    if not h_px:
        return 0.0

    cible_h = max(64, int(round(h * dpi)))
    if h_px > cible_h * 1.5:
        os.makedirs(DOSSIER_LOGOS, exist_ok=True)
        allegé = os.path.join(DOSSIER_LOGOS,
                              os.path.splitext(os.path.basename(chemin))[0] + ".png")
        if not os.path.exists(allegé) or os.path.getmtime(allegé) < os.path.getmtime(chemin):
            with _Image.open(chemin) as im:
                im = im.convert("RGBA") if im.mode in ("RGBA", "LA", "P") else im.convert("RGB")
                im.thumbnail((cible_h * 4, cible_h), _Image.LANCZOS)
                im.save(allegé, optimize=True)
            print("  logo allégé :", os.path.basename(allegé))
        chemin = allegé
        with _Image.open(chemin) as im:
            w_px, h_px = im.size

    w = h * w_px / float(h_px)
    slide.shapes.add_picture(chemin, Inches(x), Inches(y), Inches(w), Inches(h))
    return w


def slide_title(prs, blank):
    s = prs.slides.add_slide(blank)

    # --- bandeau blanc : les deux logos de l'institution -------------------
    # Le logo de la FLE est un JPEG a fond blanc : le poser sur un bandeau
    # blanc evite tout rectangle gris visible sur le fond marine.
    H_BANDE = 1.50
    add_shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, H_BANDE, fill=WHITE)
    add_shape(s, MSO_SHAPE.RECTANGLE, 0, H_BANDE, 13.333, 0.045, fill=AMBER)

    h_logo = 1.06
    y_logo = (H_BANDE - h_logo) / 2
    x = 0.80
    x += _logo(s, LOGO_UNIV, x, y_logo, h_logo) + 0.34
    x += _logo(s, LOGO_FLE, x, y_logo, h_logo) + 0.46

    add_shape(s, MSO_SHAPE.RECTANGLE, x, 0.48, 0.012, 0.54, fill=LINE_C)

    tf = add_textbox(s, x + 0.30, 0.42, 13.333 - x - 1.10, 0.68,
                     anchor=MSO_ANCHOR.MIDDLE)
    add_paragraph(tf, [("Université de Mostaganem  ·  Faculté des Langues Étrangères",
                        {"size": 12.5, "bold": True, "color": NAVY})], space_after=2)
    add_paragraph(tf, [("2ᵉ année PEP  ·  Année universitaire 2026 – 2027",
                        {"size": 11, "color": GREY})], space_after=0)

    # --- corps marine ------------------------------------------------------
    add_shape(s, MSO_SHAPE.RECTANGLE, 0, H_BANDE + 0.045, 13.333,
              7.5 - H_BANDE - 0.045, fill=NAVY)
    add_shape(s, MSO_SHAPE.RECTANGLE, 0, H_BANDE + 0.045, 0.18,
              7.5 - H_BANDE - 0.045, fill=AMBER)

    add_shape(s, MSO_SHAPE.RECTANGLE, 1.0, 2.42, 1.6, 0.07, fill=AMBER)

    tf = add_textbox(s, 1.0, 2.72, 11.3, 1.1)
    add_paragraph(tf, [("Introduction à l'Intelligence Artificielle",
                        {"size": 40, "bold": True, "color": WHITE})],
                  line_spacing=1.0, space_after=0)

    if BILINGUE:
        tf = add_textbox(s, 1.0, 3.92, 11.3, 0.5)
        add_paragraph(tf, [("Introduction to Artificial Intelligence",
                            {"size": 18, "italic": True, "color": RGBColor(0x9E, 0xBA, 0xDA)})],
                      space_after=0)
    else:
        tf = add_textbox(s, 1.0, 3.96, 10.6, 0.5)
        add_paragraph(tf, [("Comprendre l'IA, ses usages et ses limites — "
                            "puis réaliser sa première IA",
                            {"size": TAILLE_TEXTE_MAX - 2.0, "italic": True,
                             "color": RGBColor(0x9E, 0xBA, 0xDA)})],
                      space_after=0)

    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 1.0, 4.72, 5.1, 0.58,
              fill=AMBER, radius=0.35)
    tf = add_textbox(s, 1.0, 4.72, 5.1, 0.58, anchor=MSO_ANCHOR.MIDDLE)
    add_paragraph(tf, [("Questions & Réponses — Cours 01",
                        {"size": 15, "bold": True, "color": NAVY})],
                  align=PP_ALIGN.CENTER, space_after=0)

    add_shape(s, MSO_SHAPE.RECTANGLE, 1.0, 5.72, 11.33, 0.012,
              fill=RGBColor(0x2C, 0x45, 0x66))

    tf = add_textbox(s, 1.0, 5.96, 11.3, 1.0)
    add_paragraph(tf, [("Dr. Madani BELACEL — MCB",
                        {"size": 16, "bold": True, "color": WHITE})], space_after=3)
    add_paragraph(tf, [("Module Intelligence Artificielle  ·  10 idées  ·  32 questions et réponses",
                        {"size": 12.5, "color": RGBColor(0xC5, 0xD6, 0xEA)})],
                  space_after=0)

    tf = add_textbox(s, 1.0, 6.86, 11.3, 0.3)
    add_paragraph(tf, [("Cours en ligne : madani-belacel.github.io/site-enseignement-belacel",
                        {"size": 10, "color": RGBColor(0x7E, 0x93, 0xAD)})],
                  space_after=0)

    s.notes_slide.notes_text_frame.text = (
        "Page de garde — Q/R IA Cours 01 : Introduction à l'Intelligence Artificielle.\n"
        "Dr. Madani BELACEL — MCB — Université de Mostaganem.\n"
        "10 idées, 32 questions-réponses (%s).\n"
        "Séance 2ème année PEP — ENS, filière Français : présentation en "
        "français seul.\n"
        "Les Q/R sont dans les notes : mode Présentateur (Affichage → Mode "
        "Présentateur) pour les lire pendant la séance."
        % ("français / anglais" if BILINGUE else "français")
    )
    return s


def slide_idea(prs, blank, num, idea, page_no):
    s = prs.slides.add_slide(blank)
    add_shape(s, MSO_SHAPE.RECTANGLE, 0, 0, 13.333, 0.26, fill=NAVY)

    # pastille
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.75, 0.58, 1.35, 0.34,
              fill=AMBER, radius=0.45)
    tf = add_textbox(s, 0.75, 0.58, 1.35, 0.34, anchor=MSO_ANCHOR.MIDDLE)
    add_paragraph(tf, [("IDÉE %d" % num, {"size": TAILLE_PASTILLE, "bold": True, "color": NAVY})],
                  align=PP_ALIGN.CENTER, space_after=0)

    # titre FR — en francais seul il peut etre plus grand
    tf = add_textbox(s, 0.75, 1.06, 11.85, 0.70)
    add_paragraph(tf, [(idea["fr_t"],
                        {"size": TAILLE_TITRE - (5 if BILINGUE else 0),
                         "bold": True, "color": NAVY})],
                  line_spacing=1.0, space_after=0)
    if BILINGUE:
        tf = add_textbox(s, 0.75, 1.68, 11.85, 0.36)
        add_paragraph(tf, [(idea["en_t"], {"size": 14.5, "italic": True, "color": GREY})],
                      space_after=0)

    add_shape(s, MSO_SHAPE.RECTANGLE, 0.75, 2.14, 11.83, 0.012, fill=LINE_C)

    # Cadre explicatif dimensionne sur le texte reel : ni trop court (le texte
    # deborde alors du cadre, comme sur l'idee 3 avec des hauteurs fixes), ni
    # trop grand (les idees courtes laissaient un vide enorme).
    flows = idea["flows"]
    fr = idea["fr_x"] + ("FR   " if BILINGUE else "")
    en = idea["en_x"] + "EN   " if BILINGUE else ""
    h_schema = 1.42 if len(flows) == 1 else 2 * 0.86 + 0.14
    budget = 6.72 - 2.24 - h_schema - 0.14      # place restante pour le cadre
    maxi_fr = TAILLE_TEXTE_MAX - 5.5 if BILINGUE else TAILLE_TEXTE_MAX

    t_fr, t_en = maxi_fr, 12.5
    while True:
        h_fr = max(0.34, hauteur_requise(fr, 11.63, t_fr))
        h_en = max(0.30, hauteur_requise(en, 11.63, t_en)) if BILINGUE else 0.0
        ecart = 0.06 if BILINGUE else 0.0
        if h_fr + h_en + ecart + 0.28 <= budget or t_fr <= 8.5:
            break
        t_fr = round(t_fr - 0.5, 1)
        t_en = round(max(8.5, t_en - 0.5), 1)

    y_fr = 2.38
    y_en = y_fr + h_fr + 0.06
    add_shape(s, MSO_SHAPE.ROUNDED_RECTANGLE, 0.55, 2.24,
              12.23, (y_en + h_en) - 2.24 + 0.14, fill=GREY_L, radius=0.05)

    tf = add_textbox(s, 0.85, y_fr, 11.63, h_fr)
    if BILINGUE:
        add_paragraph(tf, [("FR   ", {"size": 10, "bold": True, "color": BLUE}),
                           (idea["fr_x"], {"size": t_fr, "color": DARK})],
                      line_spacing=1.15, space_after=0)
    else:
        add_paragraph(tf, [(idea["fr_x"], {"size": t_fr, "color": DARK})],
                      line_spacing=1.15, space_after=0)

    if BILINGUE:
        tf = add_textbox(s, 0.85, y_en, 11.63, h_en)
        add_paragraph(tf, [("EN   ", {"size": 10, "bold": True, "color": BLUE}),
                           (idea["en_x"], {"size": t_en, "italic": True, "color": GREY})],
                      line_spacing=1.15, space_after=0)

    # schema centre dans la place restante
    y_schema = (y_en + h_en) + 0.28
    if len(flows) == 1:
        y_schema += max(0.0, (6.72 - y_schema - h_schema) / 2)
        draw_flow(s, flows[0], 0.75, y_schema, 11.83, h_schema)
    else:
        y_schema += max(0.0, (6.72 - y_schema - h_schema) / 2)
        draw_flow(s, flows[0], 0.75, y_schema, 11.83, 0.86)
        draw_flow(s, flows[1], 0.75, y_schema + 1.00, 11.83, 0.86)

    add_footer(s, page_no)

    # ---- notes : Q/R, dans les langues affichees ----
    lines = ["IDÉE %d — %s" % (num, idea["fr_t"])]
    if BILINGUE:
        lines.append("IDEA %d — %s" % (num, idea["en_t"]))
    lines += ["", "── QUESTIONS & RÉPONSES (FR) ──"]
    for q, r in idea["qa_fr"]:
        lines += [q, r, ""]
    if BILINGUE:
        lines.append("── QUESTIONS & ANSWERS (EN) ──")
        for q, r in idea["qa_en"]:
            lines += [q, r, ""]
    s.notes_slide.notes_text_frame.text = "\n".join(lines).strip()
    return s


def main():
    # Le cours HTML est la source de verite : on y prend titres, explications
    # et Q/R. Les textes ecrits plus haut dans ce fichier ne servent que de
    # secours si le HTML est absent ou illisible.
    source = "textes du script (secours)"
    try:
        cours = lire_cours_html(COURS_HTML)
        if len(cours) != len(IDEES):
            raise ValueError("%d idees dans le HTML, %d dans le script"
                             % (len(cours), len(IDEES)))
        for idee, lue in zip(IDEES, cours):
            for cle in ("fr_t", "en_t", "fr_x", "en_x", "qa_fr", "qa_en"):
                idee[cle] = lue[cle]
        source = "cours HTML (%s)" % os.path.basename(COURS_HTML)
    except Exception as e:
        print("ATTENTION : HTML non utilise (%s)" % e)
    print("Contenu issu de :", source)

    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]

    slide_title(prs, blank)
    for i, idea in enumerate(IDEES, start=1):
        slide_idea(prs, blank, i, idea, page_no=i + 1)

    out_dir = os.path.join(DOSSIER, "Presentations_PPTX")
    os.makedirs(out_dir, exist_ok=True)
    out = os.path.join(out_dir, "QR_IA_01_Introduction_IA.pptx")
    prs.save(out)
    print("OK ->", out)
    print("Diapositives :", len(prs.slides.__iter__.__self__._sldIdLst))
    print("Q/R FR :", sum(len(i["qa_fr"]) for i in IDEES),
          "| Q/R EN :", sum(len(i["qa_en"]) for i in IDEES))
    print("Taille   : %.1f Ko" % (os.path.getsize(out) / 1024))


if __name__ == "__main__":
    main()
