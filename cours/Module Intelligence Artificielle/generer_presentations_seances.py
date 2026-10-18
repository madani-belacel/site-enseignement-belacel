#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Génère les présentations PPTX + notes PDF des 15 séances (Cours 01 → 15).

Chaque séance seances/seance-NN.html produit :
  Presentations_PPTX/Seance_NN_<slug>.pptx     (projection en classe)
  Presentations_PPTX/Seance_NN_<slug>_notes.pdf (fiche A4 imprimable)

Mise en forme identique au reste du module (palette bleu marine/or,
format 16:9). Les Cours 01 et 02 utilisent leurs présentations déjà
importées — ce script ne les régénère pas (skip par défaut).

Usage :  python3 generer_presentations_seances.py [--tous]
"""
import os
import re
import subprocess
import sys
import tempfile
import unicodedata

from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN

DOSSIER = os.path.dirname(os.path.abspath(__file__))
SEANCES = os.path.join(DOSSIER, "seances")
SORTIE = os.path.join(DOSSIER, "Presentations_PPTX")

NAVY = RGBColor(0x12, 0x2A, 0x4F)
BLUE = RGBColor(0x1F, 0x5C, 0x99)
AMBER = RGBColor(0xC2, 0x6A, 0x0A)
GREY = RGBColor(0x5B, 0x68, 0x79)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)

NOTE_CSS = """
@page { size: A4; margin: 15mm; }
body { font-family: "DejaVu Serif", Georgia, serif; font-size: 10.5pt;
       line-height: 1.45; color: #1b2433; }
h1 { font-size: 16pt; color: #122a4f; margin: 0 0 2mm; }
h2 { font-size: 12pt; color: #122a4f; margin: 5mm 0 2mm;
     border-left: 4px solid #c26a0a; padding-left: 3mm; }
.sous { font-size: 9pt; color: #5b6879; margin: 0 0 4mm; }
li { margin: 1.5mm 0; }
"""


def texte(html):
    """HTML → texte simple."""
    t = re.sub(r'<[^>]+>', ' ', html)
    t = t.replace('&nbsp;', ' ').replace('&amp;', '&').replace("'", "'")
    return re.sub(r'\s+', ' ', t).strip()


def section(s, nom):
    m = re.search(r'<h2[^>]*>[^<]*' + re.escape(nom) + r'</h2>(.*?)(?=<h2|$)', s, re.S)
    return m.group(1) if m else ''


def items(html):
    return [texte(li) for li in re.findall(r'<li>(.*?)</li>', html, re.S)]


def slugify(t):
    """Slug ASCII (sans accents) : noms de fichiers universels."""
    t = unicodedata.normalize('NFKD', t)
    t = ''.join(c for c in t if not unicodedata.combining(c))
    t = re.sub(r"[^0-9A-Za-z]+", "_", t).strip("_")
    return t[:45]


def slide_titre(prs, kicker, titre, sous=None):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    box = s.shapes.add_textbox(Inches(0.8), Inches(0.55), Inches(11.7), Inches(1.2))
    p = box.text_frame.paragraphs[0]
    p.text = kicker
    p.font.size = Pt(20)
    p.font.color.rgb = AMBER
    p.font.bold = True
    tf = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(4.6))
    tf.text_frame.word_wrap = True
    for i, part in enumerate(titre.split('\n')):
        p = tf.text_frame.paragraphs[0] if i == 0 else tf.text_frame.add_paragraph()
        p.text = part
        p.font.size = Pt(30 if i == 0 else 22)
        p.font.bold = (i == 0)
        p.font.color.rgb = NAVY
    if sous:
        p = tf.text_frame.add_paragraph()
        p.text = sous
        p.font.size = Pt(16)
        p.font.color.rgb = GREY
    return s


def slide_puces(prs, kicker, titre, puces):
    s = slide_titre(prs, kicker, titre)
    tf = s.shapes.add_textbox(Inches(0.8), Inches(1.6), Inches(11.7), Inches(5.2))
    tf.text_frame.word_wrap = True
    for i, item in enumerate(puces[:8]):
        p = tf.text_frame.paragraphs[0] if i == 0 else tf.text_frame.add_paragraph()
        p.text = "•  " + item
        p.font.size = Pt(20)
        p.font.color.rgb = NAVY
        p.space_after = Pt(10)
    return s


def generer(n):
    src = os.path.join(SEANCES, "seance-%02d.html" % n)
    if not os.path.exists(src):
        return None
    s = open(src, encoding="utf-8").read()
    h1 = re.search(r'<h1[^>]*>(.*?)</h1>', s, re.S)
    titre = texte(h1.group(1)) if h1 else "Séance %02d" % n
    une_phrase = texte(section(s, "En une phrase"))
    idees = items(section(s, "Les 3 idées essentielles"))
    pas_a_pas = items(section(s, "Comment faire"))
    retenir = items(section(s, "À retenir"))

    slug = slugify(titre.split("—", 1)[-1] if "—" in titre else titre)
    base = "Seance_%02d_%s" % (n, slug)

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(13.333), Inches(7.5)
    slide_titre(prs, "Module Intelligence Artificielle · PEP 2ème année",
                titre, "Dr. Madani BELACEL — Université de Mostaganem")
    if une_phrase:
        slide_puces(prs, "🎯 En une phrase", "", [une_phrase])
    if idees:
        slide_puces(prs, "🔑 Les 3 idées essentielles", "", idees)
    if pas_a_pas:
        slide_puces(prs, "📋 Comment faire — pas à pas", "", pas_a_pas)
    if retenir:
        slide_puces(prs, "🧠 À retenir", "", retenir)
    pptx_path = os.path.join(SORTIE, base + ".pptx")
    prs.save(pptx_path)

    # ── notes PDF imprimables ────────────────────────────────────────
    body = ['<h1>%s</h1>' % titre,
            '<p class="sous">Module Intelligence Artificielle — PEP 2ème année · '
            'Séance %02d / 15 · Dr. Madani BELACEL</p>' % n]
    if une_phrase:
        body.append('<h2>🎯 En une phrase</h2><p>%s</p>' % une_phrase)
    for label, liste in (("🔑 Les 3 idées essentielles", idees),
                         ("📋 Comment faire — pas à pas", pas_a_pas),
                         ("🧠 À retenir", retenir)):
        if liste:
            body.append('<h2>%s</h2><ul>%s</ul>' % (
                label, "".join("<li>%s</li>" % x for x in liste)))
    html = ("<html><head><meta charset='utf-8'><style>%s</style></head>"
            "<body>%s</body></html>" % (NOTE_CSS, "".join(body)))
    with tempfile.NamedTemporaryFile("w", suffix=".html", delete=False,
                                     encoding="utf-8") as fh:
        fh.write(html)
        tmp = fh.name
    pdf_path = os.path.join(SORTIE, base + "_notes.pdf")
    subprocess.run(["soffice", "--headless", "--convert-to", "pdf",
                    "--outdir", SORTIE, tmp], capture_output=True, timeout=180)
    gen = os.path.join(SORTIE, os.path.basename(tmp).replace(".html", ".pdf"))
    if os.path.exists(gen):
        os.replace(gen, pdf_path)
    os.unlink(tmp)
    return os.path.basename(pptx_path), os.path.basename(pdf_path)


def main():
    tous = "--tous" in sys.argv
    os.makedirs(SORTIE, exist_ok=True)
    for n in range(1, 16):
        if n in (1, 2) and not tous:     # présentations déjà importées
            continue
        r = generer(n)
        if r:
            print("✔ Cours %02d → %s + %s" % (n, r[0], r[1]))


if __name__ == "__main__":
    main()
